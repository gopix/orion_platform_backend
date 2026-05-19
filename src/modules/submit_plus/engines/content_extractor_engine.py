import os
from pathlib import Path
from src.modules.submit_plus.engines.base_engine import BaseEngine
from src.core.logger import get_logger

logger = get_logger(__name__)

class ContentExtractorEngine(BaseEngine):
    """
    Stateless engine for extracting text content from manuscript files.
    Handles both Word documents (.docx) and PDF files.
    
    This engine has NO database dependencies - it only works with file paths and content.
    """

    def extract_from_file(self, file_path: str) -> dict:
        """
        Extract content from a file.
        
        Args:
            file_path: Absolute path to the file
            
        Returns:
            Dictionary with extracted content and metadata
        """
        
        # -------------------------
        # Validation
        # -------------------------
        if not file_path:
            return {
                "content": "",
                "status": "error",
                "message": "No file path provided",
                "error_details": "file_path is required"
            }
        
        if not os.path.exists(file_path):
            return {
                "content": "",
                "status": "error",
                "message": f"File not found: {file_path}",
                "error_details": "The file does not exist"
            }
        
        file_extension = os.path.splitext(file_path)[1].lower()
        
        # -------------------------
        # Extract by file type
        # -------------------------
        try:
            if file_extension == ".docx":
                return self._extract_from_docx(file_path)
            elif file_extension == ".pdf":
                return self._extract_from_pdf(file_path)
            else:
                return {
                    "content": "",
                    "status": "error",
                    "message": f"Unsupported file format: {file_extension}",
                    "error_details": "Only .docx and .pdf files are supported"
                }
        except Exception as e:
            logger.error(f"Error extracting content from {file_path}: {str(e)}", exc_info=True)
            return {
                "content": "",
                "status": "error",
                "message": f"Failed to extract content: {str(e)}",
                "error_details": str(e)
            }

    def run(self, data: dict):
        """
        Placeholder for base engine interface.
        The actual extraction should use extract_from_file() method.
        """
        # This method is kept for compatibility with BaseEngine
        # In the new architecture, service layer should call extract_from_file() directly
        file_path = data.get("file_path")
        if file_path:
            return self.extract_from_file(file_path)
        return {
            "content": "",
            "status": "error",
            "message": "file_path not provided in data dict"
        }

    def _extract_from_docx(self, file_path: str) -> dict:
        """Extract text content from a Word document (.docx)"""
        try:
            from docx import Document
            
            doc = Document(file_path)
            
            # Extract text from all paragraphs
            content_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    content_parts.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text.strip())
                    content_parts.append(" | ".join(row_text))
            
            content = "\n".join(content_parts)
            
            return {
                "content": content,
                "status": "success",
                "file_type": "docx",
                "file_path": file_path,
                "word_count": len(content.split()),
                "character_count": len(content),
                "message": f"Successfully extracted {len(content_parts)} sections from DOCX"
            }
            
        except ImportError:
            logger.error("python-docx is not installed. Install it using: pip install python-docx")
            return {
                "content": "",
                "status": "error",
                "message": "python-docx library is not installed",
                "error_details": "Install python-docx to process DOCX files"
            }

    def _extract_from_pdf(self, file_path: str) -> dict:
        """Extract text content from a PDF document"""
        try:
            import PyPDF2
            
            content_parts = []
            
            with open(file_path, 'rb') as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                num_pages = len(pdf_reader.pages)
                
                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        content_parts.append(text)
            
            content = "\n".join(content_parts)
            
            return {
                "content": content,
                "status": "success",
                "file_type": "pdf",
                "file_path": file_path,
                "page_count": num_pages,
                "word_count": len(content.split()),
                "character_count": len(content),
                "message": f"Successfully extracted content from {num_pages} pages of PDF"
            }
            
        except ImportError:
            logger.error("PyPDF2 is not installed. Install it using: pip install PyPDF2")
            return {
                "content": "",
                "status": "error",
                "message": "PyPDF2 library is not installed",
                "error_details": "Install PyPDF2 to process PDF files"
            }
